# [IMPORT] Libraries
import os
import io
import sys
import json
from datetime import datetime
from openpyxl import load_workbook  # pyright: ignore[reportMissingModuleSource]
from openpyxl.cell.cell import MergedCell  # pyright: ignore[reportMissingModuleSource]
from openpyxl.cell.cell import Cell  # pyright: ignore[reportMissingModuleSource]

# [IMPORT] Utilities
from sf.utils.normalization import safe

sys.stdin = io.TextIOWrapper(sys.stdin.buffer, encoding="utf-8")

GRADE_YEAR_MAP = {
    "7": "Year I",
    "8": "Year II",
    "9": "Year III",
    "10": "Year IV",
}

try:
    # =========================
    # LOAD PAYLOAD
    # =========================
    payload = json.loads(sys.stdin.read())

    students_list = payload.get("students", [])
    adviser = payload.get("adviser", {})
    section = payload.get("section", {})
    school = payload.get("school", {})
    paths = payload.get("paths", {})

    TEMPLATE_PATH = paths.get("templatePath")
    OUTPUT_PATH = paths.get("outputPath")

    if not TEMPLATE_PATH or not OUTPUT_PATH:
        raise ValueError("Missing templatePath or outputPath")

    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)

    # =========================
    # LOAD TEMPLATE
    # =========================
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active

    # =========================
    # SAFE CELL WRITER
    # =========================
    def get_safe_cell(ws, row, col) -> Cell | None:
        cell = ws.cell(row=row, column=col)

        if isinstance(cell, MergedCell):
            for merged in ws.merged_cells.ranges:
                if cell.coordinate in merged:
                    return ws.cell(row=merged.min_row, column=merged.min_col)
            return None

        return cell

    def write_cell(row, col, value, as_text=False):
        if value is None or value == "":
            return

        cell = get_safe_cell(ws, row, col)
        if cell is None:
            return

        cell.value = str(value) if as_text else value

    # =========================
    # HEADER VALUES
    # =========================
    today = datetime.today()
    school_year = f"{today.year} - {today.year + 1}"
    report_month = today.strftime("%B")

    # =========================
    # FORMAT FIXES
    # =========================

    # Grade Level format: "Grade 7 (Year I)"
    raw_grade = str(section.get("gradeLevel", "")).replace("Grade", "").strip()

    year_label = GRADE_YEAR_MAP.get(raw_grade, "")

    grade_level_formatted = (
        f"Grade {raw_grade} ({year_label})"
        if raw_grade and year_label
        else ""
    )

    # Section must be uppercase
    section_name = safe(section.get("name", "")).upper()

    # =========================
    # SF2 HEADER LAYOUT
    # =========================

    # Row 3
    write_cell(3, 6, school.get("School ID", ""))       # F3 (F-I merged)
    write_cell(3, 13, school_year)                      # M3 (M-R merged)
    write_cell(3, 27, report_month)                     # AA3 (AA-AG merged)

    # Row 4
    write_cell(4, 6, school.get("School Name", ""))     # F4 (F-R merged)
    write_cell(4, 27, grade_level_formatted)            # AA4 (fixed format)
    write_cell(4, 39, section_name)                     # AM4 (UPPERCASE)

    # =========================
    # SAVE FILE
    # =========================
    wb.save(OUTPUT_PATH)

    print(f"[SUCCESS] SF2 saved to {OUTPUT_PATH}")

except Exception as e:
    print(f"[ERROR] {e}", file=sys.stderr)
    sys.exit(1)