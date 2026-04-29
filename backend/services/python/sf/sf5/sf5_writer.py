# [IMPORT] Libraries
import os
import io
import sys
import json
from datetime import datetime
from openpyxl import load_workbook  # pyright: ignore[reportMissingModuleSource]
from openpyxl.cell.cell import MergedCell  # pyright: ignore[reportMissingModuleSource]
from openpyxl.cell.cell import Cell  # pyright: ignore[reportMissingModuleSource]

from sf.utils.normalization import safe

sys.stdin = io.TextIOWrapper(sys.stdin.buffer, encoding="utf-8")

try:
    # =========================
    # LOAD PAYLOAD
    # =========================
    payload = json.loads(sys.stdin.read())

    school = payload.get("school", {})
    section = payload.get("section", {})
    adviser = payload.get("adviser", {})
    paths = payload.get("paths", {})

    TEMPLATE_PATH = paths.get("templatePath")
    OUTPUT_PATH = paths.get("outputPath")

    if not TEMPLATE_PATH or not OUTPUT_PATH:
        raise ValueError("Missing templatePath or outputPath")

    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)

    # =========================
    # LOAD TEMPLATE
    # =========================
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active

    # =========================
    # SAFE WRITER
    # =========================
    def get_safe_cell(ws, row, col) -> Cell | None:
        cell = ws.cell(row=row, column=col)

        if isinstance(cell, MergedCell):
            for merged in ws.merged_cells.ranges:
                if cell.coordinate in merged:
                    return ws.cell(row=merged.min_row, column=merged.min_col)
            return None

        return cell

    def write_cell(row, col, value):
        if value is None or value == "":
            return

        cell = ws.cell(row=row, column=col)

        print(f"[WRITE] R{row} C{col} = {value}")  # DEBUG TRACE

        if isinstance(cell, MergedCell):
            for merged in ws.merged_cells.ranges:
                if cell.coordinate in merged:
                    anchor = ws.cell(row=merged.min_row, column=merged.min_col)
                    print(f"[MERGE FIX] {cell.coordinate} -> {anchor.coordinate}")
                    anchor.value = value
                    return
            return

        cell.value = value

    # =========================
    # NAME FORMAT (FIXED)
    # FIRST MIDDLE LAST (NO COMMAS)
    # =========================
    def full_name(person):
        first = safe(person.get("firstName", person.get("First Name", ""))).upper()
        middle = safe(person.get("middleName", person.get("Middle Name", ""))).upper()
        last = safe(person.get("lastName", person.get("Last Name", ""))).upper()

        return " ".join([p for p in [first, middle, last] if p]).strip()

    # =========================
    # HEADER MAPPINGS
    # =========================

    # Row 3
    write_cell(3, 3, school.get("Region", ""))        # C3
    write_cell(3, 5, school.get("Division", ""))      # E-H (merged)
    write_cell(3, 10, school.get("District", ""))     # J-L (merged)

    # Row 5
    write_cell(5, 3, school.get("School ID", ""))     # C-D
    today = datetime.today()
    school_year = f"{today.year} - {today.year + 1}"

    write_cell(5, 7, school_year)
    curriculum = safe(section.get("curriculum", ""))

    write_cell(5, 10, curriculum)  # J
    write_cell(5, 11, curriculum)  # K
    write_cell(5, 12, curriculum)  # L

    # Row 7
    write_cell(7, 3, school.get("School Name", ""))   # C-H
    write_cell(7, 10, section.get("gradeLevel", ""))  # J
    write_cell(7, 13, safe(section.get("name", "")).upper())  # M-O

    # =========================
    # FOOTER FIELDS
    # =========================

    # Adviser (Row 35, Col L-O)
    write_cell(35, 12, full_name(adviser))  # L
    write_cell(35, 13, full_name(adviser))  # M
    write_cell(35, 14, full_name(adviser))  # N
    write_cell(35, 15, full_name(adviser))  # O

    # =========================
    # SCHOOL HEAD (FIXED: STRING, NOT OBJECT)
    # =========================

    school_head = safe(school.get("School Head", "")).upper()

    write_cell(40, 12, school_head)
    write_cell(40, 13, school_head)
    write_cell(40, 14, school_head)
    write_cell(40, 15, school_head)

    # =========================
    # STUDENT LIST CONFIG
    # =========================
    START_ROW = 12

    COL_NAME = 2  # Column B

    def format_name_last_first_middle(s):
        last = safe(s.get("Last Name", s.get("lastName", ""))).upper()
        first = safe(s.get("First Name", s.get("firstName", ""))).upper()
        middle = safe(s.get("Middle Name", s.get("middleName", ""))).upper()

        return f"{last}, {first}, {middle}".strip(", ").replace(", ,", ",")


    def write_student_list(data_list, start_row):
        row = start_row

        for s in data_list:
            write_cell(row, COL_NAME, format_name_last_first_middle(s))
            row += 1

        return row
  
    def next_row(r):
      return r + 1


    # =========================
    # SPLIT STUDENTS
    # =========================
    males = [s for s in payload.get("students", []) if safe(s.get("Sex")).upper() == "MALE"]
    females = [s for s in payload.get("students", []) if safe(s.get("Sex")).upper() == "FEMALE"]

    totalMaleCount = len(males)
    totalFemaleCount = len(females)
    totalCombinedCount = totalMaleCount + totalFemaleCount

    current_row = START_ROW

    # =========================
    # MALES
    # =========================
    current_row = write_student_list(males, current_row)

    write_cell(current_row, COL_NAME, f"Total Male: {totalMaleCount}")
    current_row = next_row(current_row)

    # =========================
    # FEMALES
    # =========================
    current_row = write_student_list(females, current_row)

    write_cell(current_row, COL_NAME, f"Total Female: {totalFemaleCount}")
    current_row = next_row(current_row)

    # =========================
    # TOTAL COMBINED
    # =========================
    write_cell(current_row, COL_NAME, f"Total Combined: {totalCombinedCount}")

    # =========================
    # SAVE FILE
    # =========================
    wb.save(OUTPUT_PATH)

    print(f"[SUCCESS] SF5 saved to {OUTPUT_PATH}")

except Exception as e:
    print(f"[ERROR] {e}", file=sys.stderr)
    sys.exit(1)