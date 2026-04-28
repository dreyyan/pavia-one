# [IMPORT] Libraries
import os
import io
import sys
import json
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook  # pyright: ignore[reportMissingModuleSource]
from openpyxl.cell.cell import MergedCell  # pyright: ignore[reportMissingModuleSource]
from openpyxl.cell.cell import Cell  # pyright: ignore[reportMissingModuleSource]

# [IMPORT] Utilities
from sf.utils.normalization import safe, is_male, is_female, format_birthdate
from sf.utils.name_parser import adviser_full_name

sys.stdin = io.TextIOWrapper(sys.stdin.buffer, encoding='utf-8')

try:
    # [LOAD PAYLOAD]
    payload = json.loads(sys.stdin.read())

    students_list = payload.get("students", [])
    adviser = payload.get("adviser", {})
    section = payload.get("section", {})
    school = payload.get("school", {})
    paths = payload.get("paths", {})

    TEMPLATE_PATH = paths.get("templatePath")
    OUTPUT_PATH = paths.get("outputPath")

    if not TEMPLATE_PATH or not OUTPUT_PATH:
        raise ValueError("Missing templatePath or outputPath")

    OUTPUT_DIR = os.path.dirname(OUTPUT_PATH)
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    adviser_name = adviser_full_name(adviser)

    # DEFAULT VALUES
    defaults = {
        "Age": "",
        "Mother Tongue": "",
        "Religion": "",

        "Grade Level": section.get("gradeLevel", "Grade 10"),
        "Section": section.get("name", "A"),

        # SCHOOL DATA (NEW FIX)
        "School ID": school.get("School ID", ""),
        "Region": school.get("Region", ""),
        "School Name": school.get("School Name", ""),
        "Division": school.get("Division", ""),

        "IP Ethnic Group": "",
        "Learning Modality": "",
        "Remarks": "",
        "Father Name": "",
        "Mother Maiden Name": "",
        "Barangay": "",
        "Municipality": "",
        "Province": ""
    }

    for s in students_list:
        for k, v in defaults.items():
            s.setdefault(k, v)

    data = pd.DataFrame(students_list).fillna("")

    # LOAD TEMPLATE
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active

    # SAFE CELL WRITER
    def get_safe_cell(ws, row, col) -> Cell | None:
        cell = ws.cell(row=row, column=col)

        if isinstance(cell, MergedCell):
            for merged in ws.merged_cells.ranges:
                if cell.coordinate in merged:
                    return ws.cell(row=merged.min_row, column=merged.min_col)
            return None

        return cell

    def write_cell(row, col, value, as_text=False):
        if value in ("", None):
            return

        cell = get_safe_cell(ws, row, col)
        if cell is None:
            return

        cell.value = str(value) if as_text else value

    # HEADER
    today = datetime.today()
    sy = f"{today.year} - {today.year + 1}"

    for c in range(20, 25):
        write_cell(4, c, sy)

    # Section Details
    grade = data["Grade Level"].iloc[0] if len(data) else "Grade 10"
    section_name = data["Section"].iloc[0].upper() if len(data) else "A"

    write_cell(4, 31, grade)
    write_cell(4, 32, grade)

    # School Details
    # Row 3
    write_cell(3, 6, school.get("School ID", ""))   # F3 (F-I merged)
    write_cell(3, 11, school.get("Region", ""))     # K3 (K-O merged)
    write_cell(3, 20, school.get("Division", ""))   # T3 (T-AF merged)

    # Row 4
    write_cell(4, 6, school.get("School Name", "")) # F4 (F-O merged)

    for c in range(39, 48):
        write_cell(4, c, section_name)
        

    # =========================
    # CONFIG
    # =========================
    START_ROW = 7

    COL = {
        "lrn": 1, "name": 3, "sex": 7, "birth": 8, "age": 10,
        "mother_tongue": 12, "ip": 14, "religion": 15,
        "barangay": 18, "municipality": 21, "province": 23,
        "father": 28, "mother": 32, "modality": 44, "remarks": 45
    }

    # =========================
    # HELPERS
    # =========================
    def format_name(n):
        n = safe(n).strip()
        if not n:
            return ""

        if "," in n:
            a, b = n.split(",", 1)
            return f"{a.strip().upper()}, {b.strip().upper()}"

        parts = n.split()
        if len(parts) == 1:
            return parts[0].upper()
        if len(parts) == 2:
            return f"{parts[1].upper()}, {parts[0].upper()}"

        return f"{parts[-1].upper()}, {parts[0].upper()} {' '.join(parts[1:-1]).upper()}"

    def full_name(s):
        return f"{safe(s.get('Last Name')).upper()}, {safe(s.get('First Name')).upper()} {safe(s.get('Middle Name')).upper()}"

    def write_student(s, row):
        sex = "M" if is_male(s.get("Sex")) else "F"

        write_cell(row, COL["lrn"], s.get("LRN", ""), True)
        write_cell(row, COL["name"], full_name(s))
        write_cell(row, COL["sex"], sex)
        write_cell(row, COL["birth"], format_birthdate(s.get("Birth Date")))
        write_cell(row, COL["age"], s.get("Age", ""))

        write_cell(row, COL["mother_tongue"], s.get("Mother Tongue", ""))
        write_cell(row, COL["ip"], safe(s.get("IP Ethnic Group")).upper())
        write_cell(row, COL["religion"], s.get("Religion", ""))

        write_cell(row, COL["barangay"], safe(s.get("Barangay")).upper())
        write_cell(row, COL["municipality"], safe(s.get("Municipality")).upper())
        write_cell(row, COL["province"], safe(s.get("Province")).upper() or "ILOILO")

        write_cell(row, COL["father"], format_name(s.get("Father Name")))
        write_cell(row, COL["mother"], format_name(s.get("Mother Maiden Name")))

        write_cell(row, COL["modality"], safe(s.get("Learning Modality")).replace("_", " ").title())
        write_cell(row, COL["remarks"], s.get("Remarks", ""))

    # =========================
    # SPLIT DATA
    # =========================
    males = data[data["Sex"].apply(is_male)].reset_index(drop=True)
    females = data[data["Sex"].apply(is_female)].reset_index(drop=True)

    male_count = len(males)
    female_count = len(females)
    grand_total = male_count + female_count

    # =========================
    # CONSTANT TEMPLATE LIMITS
    # =========================
    START_ROW = 7
    END_ROW = 68

    current = START_ROW

    # =========================
    # VALIDATE SPACE
    # =========================
    students_total = len(males) + len(females)

    MAX_DATA_ROWS = END_ROW - START_ROW + 1

    # +3 for totals (male, female, grand total)
    required_rows = students_total + 3

    if required_rows > MAX_DATA_ROWS:
        raise ValueError(
            f"Too many students ({students_total}). "
            f"Template only supports {MAX_DATA_ROWS - 3} students (reserve space for totals)."
        )

    # =========================
    # DATA AREA CONFIG
    # =========================
    DATA_START = 7
    DATA_END = 68

    current = DATA_START

    students_total = len(males) + len(females)

    # +3 total rows (male + female + grand total)
    required_rows = students_total + 3

    max_rows = DATA_END - DATA_START + 1

    if required_rows > max_rows:
        raise ValueError(
            f"Too many students ({students_total}). "
            f"Only {max_rows - 3} allowed."
        )

    # =========================
    # WRITE MALES
    # =========================
    for _, s in males.iterrows():
        write_student(s, current)
        current += 1

    write_cell(current, 1, male_count)
    write_cell(current, 2, male_count)
    write_cell(current, 3, "<=== TOTAL MALE")
    current += 1

    # =========================
    # WRITE FEMALES
    # =========================
    for _, s in females.iterrows():
        write_student(s, current)
        current += 1

    write_cell(current, 1, female_count)
    write_cell(current, 2, female_count)
    write_cell(current, 3, "<=== TOTAL FEMALE")
    current += 1

    # =========================
    # GRAND TOTAL
    # =========================
    write_cell(current, 1, grand_total)
    write_cell(current, 2, grand_total)
    write_cell(current, 3, "<=== TOTAL COMBINED")

    # =========================
    # TOTALS + ADVISER (MERGED CELLS SAFE)
    # =========================

    # Total Male (Row 71–73, Col X–Z → X = 24)
    write_cell(71, 24, male_count)

    # Total Female (Row 74–75, Col X–Z)
    write_cell(74, 24, female_count)

    # Total Combined (Row 76–77, Col X–Z)
    write_cell(76, 24, grand_total)

    # Adviser Name (Row 71–73, Col AE–AK → AE = 31)
    write_cell(
        71,
        31,
        adviser_full_name(adviser)
    )

    # =========================
    # SIMULATED ROW DELETION (NO SHIFTING, NO STRUCTURE CHANGE)
    # =========================

    def wipe_row(row):
        for c in range(1, 50):
            cell = ws.cell(row=row, column=c)

            # skip merged "slave" cells safely
            if isinstance(cell, MergedCell):
                continue

            cell.value = None


    last_used_row = current

    # "delete" remaining rows in SF1 data area
    for r in range(last_used_row + 1, DATA_END + 1):
        wipe_row(r)

    # =========================
    # SAVE FILE
    # =========================
    wb.save(OUTPUT_PATH)

    print(f"[SUCCESS] Saved to {OUTPUT_PATH}")
    print(f"Male: {male_count}, Female: {female_count}, Total: {grand_total}")

except Exception as e:
    print(f"[ERROR] {e}", file=sys.stderr)
    sys.exit(1)